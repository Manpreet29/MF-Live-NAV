"""
Parses monthly portfolio disclosure Excel files from Indian AMCs.

Supported formats (auto-detected):
  - HDFC Mutual Fund   (sheet: HDFCEQ, header: 'Name Of the Instrument')
  - ICICI Prudential   (sheet: MULTI,  header: 'Company/Issuer/Instrument Name')
  - Any AMC that follows a similar tabular layout (best-effort)

Key format differences handled:
  HDFC:  Name=col3, ISIN=col1, Qty=col5, MktVal=col6, Pct=col7  (% as number e.g. 8.69)
  ICICI: Name=col1, ISIN=col2, Qty=col5, MktVal=col6, Pct=col7  (% as decimal e.g. 0.0869)

Market Value is in Rs. LACS in both formats:
  disclosed_price = market_value_lacs * 100_000 / quantity
"""

import re
import logging
from io import BytesIO
from typing import Optional
from openpyxl import load_workbook

logger = logging.getLogger(__name__)

# ------------------------------------------------------------------ #
# Detection constants                                                  #
# ------------------------------------------------------------------ #

# All known header signals across AMC formats
HEADER_SIGNALS = [
    "name of the instrument",           # HDFC
    "company/issuer/instrument name",   # ICICI
    "name of the issuer",               # Some other AMCs
    "instrument name",                  # Generic fallback
]

# Valid 12-char ISIN (2 alpha + 10 alphanumeric)
ISIN_RE = re.compile(r"^[A-Z]{2}[A-Z0-9]{10}$")

# Stop parsing when any cell in a row exactly matches one of these
# (case-insensitive). These signal the start of non-equity sections.
HARD_STOP_SIGNALS = frozenset([
    "debt instruments",
    "money market instruments",
    "government securities",
    "grand total",
    "total net assets",
    "net current assets",
    "treps",
    "certificate of deposits",
    "commercial papers",
    "treasury bills",
    "compulsory convertible debenture",
    "non-convertible debentures / bonds",
    "zero coupon bonds / deep discount bonds",
    "term deposits",
    "securitized debt instruments",
    "non convertible preference shares",
    "privately placed/unlisted",
    "privately placed/unlisted",
    "units of mutual fund",
    "units of an alternative investment fund",
    "others",
    "gold",
    "interest rate swaps",
])

# Rows that are section sub-headers — skip but don't stop
SECTION_SKIP_SIGNALS = frozenset([
    "equity & equity related instruments",
    "listed / awaiting listing on stock exchanges",
    "units of real estate investment trust (reit)",
    "units of infrastructure investment trust",
    "unlisted",
    "listed",
    "equity",
])

# ------------------------------------------------------------------ #
# Data class                                                           #
# ------------------------------------------------------------------ #

class ParsedHolding:
    def __init__(self, name, isin, industry, quantity,
                 market_value_lacs, disclosed_price, pct_to_nav):
        self.name              = name
        self.isin              = isin
        self.industry          = industry
        self.quantity          = quantity
        self.market_value_lacs = market_value_lacs
        self.disclosed_price   = disclosed_price
        self.pct_to_nav        = pct_to_nav


class HoldingsParseError(Exception):
    pass


# ------------------------------------------------------------------ #
# Public API                                                           #
# ------------------------------------------------------------------ #

def parse_holdings(file_bytes: bytes) -> tuple[list[ParsedHolding], str]:
    """
    Parse an AMC monthly portfolio disclosure Excel.
    Auto-detects the AMC format.
    Returns (holdings, portfolio_date_string).
    """
    try:
        wb = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as e:
        raise HoldingsParseError(f"Cannot open Excel workbook: {e}")

    sheet          = _find_sheet(wb)
    rows           = list(sheet.iter_rows(values_only=True))
    wb.close()

    if not rows:
        raise HoldingsParseError("Workbook appears to be empty.")

    portfolio_date      = _extract_date(rows)
    header_idx, col_map = _find_header(rows)
    holdings            = _extract(rows, header_idx, col_map)

    if not holdings:
        raise HoldingsParseError(
            "No valid equity holdings found. "
            "Please verify this is an AMC monthly portfolio disclosure file."
        )

    logger.info(
        "Parsed %d holdings from '%s' | date: %s",
        len(holdings), sheet.title, portfolio_date,
    )
    return holdings, portfolio_date


# ------------------------------------------------------------------ #
# Internal helpers                                                     #
# ------------------------------------------------------------------ #

def _find_sheet(wb):
    """
    Find the main equity holdings sheet.
    Priority: known HDFC/ICICI sheet names, then first sheet.
    """
    preferred = ["HDFCEQ", "hdfceq", "MULTI", "Sheet1", "Portfolio", "Holdings"]
    for name in preferred:
        if name in wb.sheetnames:
            logger.debug("Using sheet: %s", name)
            return wb[name]
    sheet = wb.worksheets[0]
    logger.warning("No preferred sheet found; using first sheet '%s'", sheet.title)
    return sheet


def _extract_date(rows: list) -> str:
    """Scan first 10 rows for 'portfolio as on' date string."""
    for row in rows[:10]:
        for cell in row:
            if cell and isinstance(cell, str) and "portfolio as on" in cell.lower():
                return cell.strip()
    return "Unknown"


def _find_header(rows: list) -> tuple[int, dict]:
    """
    Find the header row by looking for any known HEADER_SIGNAL.
    Returns (header_row_index, column_map).
    """
    for idx, row in enumerate(rows[:30]):
        row_lower = [str(c).strip().lower() if c else "" for c in row]
        if any(sig in " ".join(row_lower) for sig in HEADER_SIGNALS):
            col_map = _build_col_map(row)
            logger.debug("Header at row %d: %s", idx, col_map)
            return idx, col_map

    raise HoldingsParseError(
        "Header row not found. "
        f"Looked for: {HEADER_SIGNALS}. "
        "This may not be a supported AMC portfolio disclosure format."
    )


def _build_col_map(header_row: tuple) -> dict:
    """
    Dynamically build column-name → index mapping from the header row.
    Works for both HDFC and ICICI column layouts.
    """
    m: dict[str, int] = {}
    for i, cell in enumerate(header_row):
        if not cell:
            continue
        s = str(cell).strip().lower()

        # Name column (either format)
        if "name of the instrument" in s or "company/issuer/instrument name" in s or "name of the issuer" in s:
            m["name"] = i
        # ISIN — must be exact 'isin' to avoid matching 'industry'
        elif s == "isin":
            m["isin"] = i
        # Industry / Rating
        elif ("industry" in s or "rating" in s) and "isin" not in s:
            m.setdefault("industry", i)
        # Quantity
        elif "quantity" in s:
            m["quantity"] = i
        # Market / Fair / Exposure value
        elif ("market" in s or "exposure" in s or "fair" in s) and "value" in s:
            m["market_value"] = i
        # % to NAV
        elif "% to nav" in s or "%to nav" in s or "% to n" in s:
            m["pct_to_nav"] = i

    required = ["name", "isin", "quantity", "market_value", "pct_to_nav"]
    missing  = [k for k in required if k not in m]
    if missing:
        raise HoldingsParseError(
            f"Could not find required columns: {missing}. "
            f"Header cells found: {[str(c) for c in header_row if c]}"
        )
    return m


def _extract(rows: list, header_idx: int, col_map: dict) -> list[ParsedHolding]:
    """
    Iterate data rows after the header, extract valid holdings.
    Stops at hard-stop section headers.
    """
    holdings: list[ParsedHolding] = []
    pct_values: list[float]       = []   # collected to detect decimal vs percentage

    for row_idx, row in enumerate(rows[header_idx + 1:], start=header_idx + 1):
        # Build a flat list of non-empty cell strings for signal detection
        cells_lower = [str(c).strip().lower() for c in row if c is not None and str(c).strip()]

        # Hard stop — we've entered a non-equity section
        if _hits_hard_stop(cells_lower):
            logger.debug("Hard stop at row %d: %s", row_idx, cells_lower[:3])
            break

        # Soft skip — section sub-header row
        if _hits_section_skip(cells_lower):
            continue

        # Must have a valid ISIN
        isin_raw = _cell(row, col_map["isin"])
        if not isin_raw:
            continue
        isin = str(isin_raw).strip().upper()
        if not ISIN_RE.match(isin):
            continue

        holding = _build_holding(row_idx, isin, row, col_map)
        if holding:
            holdings.append(holding)
            pct_values.append(holding.pct_to_nav)

    # Auto-fix: if all pct_to_nav values are < 2.0, they're stored as decimals
    # (ICICI stores 0.0574 instead of 5.74). Multiply by 100 to normalise.
    if pct_values and max(pct_values) < 2.0:
        logger.info(
            "Detected decimal % to NAV format (max=%.4f). Multiplying by 100.",
            max(pct_values),
        )
        for h in holdings:
            h.pct_to_nav = round(h.pct_to_nav * 100, 4)

    return holdings


def _hits_hard_stop(cells_lower: list[str]) -> bool:
    for cell in cells_lower:
        if cell in HARD_STOP_SIGNALS:
            return True
        # Partial match for compound labels
        for sig in ("debt instruments", "money market", "government securities",
                    "grand total", "total net assets", "net current assets",
                    "compulsory convertible", "non-convertible debentures",
                    "zero coupon bonds", "interest rate swaps",
                    "units of mutual fund", "units of an alternative"):
            if sig in cell:
                return True
    return False


def _hits_section_skip(cells_lower: list[str]) -> bool:
    for cell in cells_lower:
        if cell in SECTION_SKIP_SIGNALS:
            return True
        for sig in ("equity & equity related", "listed / awaiting listing",
                    "units of real estate", "units of infrastructure"):
            if sig in cell:
                return True
    return False


def _cell(row: tuple, idx: int) -> Optional[object]:
    if idx < 0 or idx >= len(row):
        return None
    return row[idx]


def _build_holding(row_idx: int, isin: str, row: tuple,
                   col_map: dict) -> Optional[ParsedHolding]:
    """Parse one data row into a ParsedHolding. Returns None on bad data."""
    name_raw  = _cell(row, col_map["name"])
    ind_raw   = _cell(row, col_map.get("industry", -1))
    qty_raw   = _cell(row, col_map["quantity"])
    val_raw   = _cell(row, col_map["market_value"])
    pct_raw   = _cell(row, col_map["pct_to_nav"])

    # Name — strip leading | (HDFC top-10 marker)
    name = str(name_raw).strip().lstrip("|").strip() if name_raw else isin
    # Skip derivative/covered-call rows that sneak through
    if "$$" in name or "(covered call)" in name.lower():
        return None

    industry = str(ind_raw).strip() if ind_raw else ""

    # Quantity
    try:
        quantity = int(float(str(qty_raw).replace(",", "").strip()))
        if quantity <= 0:
            raise ValueError
    except Exception:
        logger.debug("Row %d: bad quantity '%s' for '%s'", row_idx, qty_raw, name)
        return None

    # Market value in Lacs
    try:
        market_value_lacs = float(str(val_raw).replace(",", "").strip())
        if market_value_lacs <= 0:
            raise ValueError
    except Exception:
        logger.debug("Row %d: bad market_value '%s' for '%s'", row_idx, val_raw, name)
        return None

    # % to NAV (may be decimal or percentage — normalised later in _extract)
    try:
        pct_s = str(pct_raw).strip()
        pct_to_nav = 0.005 if pct_s in ("@", "", "None") else float(pct_s.replace(",", ""))
    except Exception:
        logger.debug("Row %d: bad pct_to_nav '%s' for '%s'", row_idx, pct_raw, name)
        return None

    # Derived share price
    disclosed_price = round(market_value_lacs * 100_000 / quantity, 4)

    return ParsedHolding(
        name=name,
        isin=isin,
        industry=industry,
        quantity=quantity,
        market_value_lacs=market_value_lacs,
        disclosed_price=disclosed_price,
        pct_to_nav=pct_to_nav,
    )